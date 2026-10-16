from __future__ import annotations

import io
from typing import Optional

import pandas as pd


def dataframe_to_csv_bytes(df: pd.DataFrame) -> bytes:
    """Return UTF-8 encoded CSV bytes for a DataFrame."""
    return df.to_csv(index=False).encode("utf-8")


def dataframe_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """
    Return a single-sheet .xlsx file for Power BI / Excel import.
    Applies auto-width columns and a header style.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    header_fill = PatternFill("solid", fgColor="2D6AB5")   # Power BI blue
    header_font = Font(bold=True, color="FFFFFF")

    # Write header
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Write rows (cap at 100k for file size)
    for row_idx, row in enumerate(df.head(100_000).itertuples(index=False), start=2):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    # Auto-fit column widths
    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def figures_to_excel_bytes(
    figures: list[tuple],  # list of (title, fig) tuples
    dashboard_data: dict,
    result_df: Optional[pd.DataFrame] = None,
) -> bytes:
    """
    Build an openpyxl workbook with:
      Sheet 1: Summary text (title, insights, KPIs, recommendations)
      Sheet 2: Data table (result_df)
      Sheet 3+: Embedded chart PNGs (one per figure)
    Returns bytes of the .xlsx file.
    """
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from utils.charts import fig_to_png_bytes
    import matplotlib.pyplot as plt

    wb = openpyxl.Workbook()

    # --- Sheet 1: Summary ---
    ws_summary = wb.active
    ws_summary.title = "Summary"

    row = 1
    title = dashboard_data.get("title", "Dashboard Report")
    ws_summary.cell(row=row, column=1, value=title)
    ws_summary.cell(row=row, column=1).font = openpyxl.styles.Font(bold=True, size=14)
    row += 2

    summary = dashboard_data.get("executive_summary", "")
    if summary:
        ws_summary.cell(row=row, column=1, value="Executive Summary")
        ws_summary.cell(row=row, column=1).font = openpyxl.styles.Font(bold=True)
        row += 1
        ws_summary.cell(row=row, column=1, value=summary)
        row += 2

    kpis = dashboard_data.get("kpis", [])
    if kpis:
        ws_summary.cell(row=row, column=1, value="Key Performance Indicators")
        ws_summary.cell(row=row, column=1).font = openpyxl.styles.Font(bold=True)
        row += 1
        ws_summary.cell(row=row, column=1, value="Metric")
        ws_summary.cell(row=row, column=2, value="Value")
        ws_summary.cell(row=row, column=3, value="Change")
        row += 1
        for kpi in kpis:
            ws_summary.cell(row=row, column=1, value=kpi.get("label", ""))
            ws_summary.cell(row=row, column=2, value=kpi.get("value", ""))
            ws_summary.cell(row=row, column=3, value=kpi.get("delta", ""))
            row += 1
        row += 1

    insights = dashboard_data.get("insights", [])
    if insights:
        ws_summary.cell(row=row, column=1, value="Key Insights")
        ws_summary.cell(row=row, column=1).font = openpyxl.styles.Font(bold=True)
        row += 1
        for insight in insights:
            ws_summary.cell(row=row, column=1, value=f"• {insight}")
            row += 1
        row += 1

    recommendations = dashboard_data.get("recommendations", [])
    if recommendations:
        ws_summary.cell(row=row, column=1, value="Recommendations")
        ws_summary.cell(row=row, column=1).font = openpyxl.styles.Font(bold=True)
        row += 1
        for rec in recommendations:
            ws_summary.cell(row=row, column=1, value=f"→ {rec}")
            row += 1

    # Auto-width column A
    ws_summary.column_dimensions["A"].width = 80

    # --- Sheet 2: Data ---
    if result_df is not None and not result_df.empty:
        ws_data = wb.create_sheet("Data")
        # Header
        for col_idx, col_name in enumerate(result_df.columns, start=1):
            cell = ws_data.cell(row=1, column=col_idx, value=col_name)
            cell.font = openpyxl.styles.Font(bold=True)
        # Rows (cap at 10k for file size)
        for row_idx, row_data in enumerate(result_df.head(10_000).itertuples(index=False), start=2):
            for col_idx, val in enumerate(row_data, start=1):
                ws_data.cell(row=row_idx, column=col_idx, value=val)

    # --- Sheet 3+: Charts ---
    for idx, (fig_title, fig) in enumerate(figures):
        ws_chart = wb.create_sheet(f"Chart {idx + 1}")
        ws_chart.cell(row=1, column=1, value=fig_title)
        ws_chart.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True, size=12)

        png_bytes = fig_to_png_bytes(fig)
        img = XLImage(io.BytesIO(png_bytes))
        img.width = 600
        img.height = 400
        ws_chart.add_image(img, "A3")

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
